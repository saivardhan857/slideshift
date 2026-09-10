"""Phase 4 — diagnostics, error sanitization, upload validation, job lifecycle.

Self-contained — no pytest.  Run:  python test_v2_phase4.py
"""
import io
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))

import main
from diagnostics import build_diagnostics, demo as diag_demo
from transfer import _dedupe_pptx

_passed = _failed = 0


def check(name, cond, detail=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK]   {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name}  {detail}")


def _pptx_bytes(n_slides=1, dense=False):
    prs = Presentation()
    for _ in range(n_slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Title"
        body = s.placeholders[1].text_frame
        body.text = "point one"
        if dense:
            for i in range(60):
                body.add_paragraph().text = f"a fairly long lecture bullet number {i} to force overflow"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _sse_events(text):
    out = []
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("data: "):
            try:
                out.append(json.loads(line[6:]))
            except Exception:
                pass
    return out


def run():
    diag_demo()  # diagnostics module self-check

    # ------------------------------------------------------------------ #
    # Error sanitization
    # ------------------------------------------------------------------ #
    leak = FileNotFoundError("[Errno 2] No such file or directory: "
                             "'/tmp/ctm_jobs/abc123/source.pptx'")
    msg = main._safe_error(leak, "job abc123: transfer")
    for bad in ("/tmp", "ctm_jobs", "Errno", "Traceback", "source.pptx", "abc123"):
        check(f"safe_error hides {bad!r}", bad not in msg, msg)
    check("safe_error still says something useful", "error occurred" in msg.lower(), msg)

    # ------------------------------------------------------------------ #
    # Upload validation (pure helpers)
    # ------------------------------------------------------------------ #
    from fastapi import HTTPException

    def _expect_http(status, fn):
        try:
            fn()
            return False, "no exception"
        except HTTPException as e:
            return e.status_code == status, f"got {e.status_code}: {e.detail}"

    ok, d = _expect_http(400, lambda: main._validate_pptx_bytes(b"", "X"))
    check("empty bytes rejected (400)", ok, d)
    ok, d = _expect_http(400, lambda: main._validate_pptx_bytes(b"not a zip at all", "X"))
    check("garbage bytes rejected (400)", ok, d)
    # a real zip that is not a pptx (no ppt/presentation.xml)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("hello.txt", "hi")
    ok, d = _expect_http(400, lambda: main._validate_pptx_bytes(zbuf.getvalue(), "X"))
    check("zip without ppt/presentation.xml rejected (400)", ok, d)
    # a genuine pptx passes
    try:
        main._validate_pptx_bytes(_pptx_bytes(1), "X")
        check("valid pptx bytes accepted", True)
    except HTTPException as e:
        check("valid pptx bytes accepted", False, e.detail)

    # ------------------------------------------------------------------ #
    # _dedupe_pptx equivalence: last occurrence of a name wins, no duplicates
    # ------------------------------------------------------------------ #
    dd = Path(tempfile.mkdtemp(prefix="p4_dedupe_"))
    zp = dd / "dup.zip"
    with zipfile.ZipFile(zp, "w") as z:
        z.writestr("a.txt", "FIRST")
        z.writestr("b.txt", "bee")
        z.writestr("a.txt", "SECOND")   # duplicate name
    _dedupe_pptx(str(zp))
    with zipfile.ZipFile(zp) as z:
        names = z.namelist()
        check("dedupe removes duplicate entries", names.count("a.txt") == 1, names)
        check("dedupe keeps the LAST occurrence", z.read("a.txt") == b"SECOND", z.read("a.txt"))
        check("dedupe leaves other entries intact", z.read("b.txt") == b"bee")

    # ------------------------------------------------------------------ #
    # Job lifecycle helper
    # ------------------------------------------------------------------ #
    import asyncio
    jd = Path(tempfile.mkdtemp(prefix="p4_job_")) / "job1"
    jd.mkdir()
    (jd / "f").write_text("x")
    asyncio.get_event_loop().run_until_complete(main._delayed_cleanup(jd, delay=0))
    check("_delayed_cleanup removes the job directory", not jd.exists())

    # ------------------------------------------------------------------ #
    # Integration via TestClient — isolate TEMP_DIR
    # ------------------------------------------------------------------ #
    from fastapi.testclient import TestClient

    real_tmp = main.TEMP_DIR
    main.TEMP_DIR = Path(tempfile.mkdtemp(prefix="p4_temp_"))
    orig_max_bytes = main.MAX_UPLOAD_BYTES
    orig_max_slides = main.MAX_SLIDES
    try:
        with TestClient(main.app) as client:
            check("GET /health -> 200", client.get("/health").status_code == 200)

            good = _pptx_bytes(1)
            tpl = _pptx_bytes(1)

            # bad extension
            r = client.post("/api/transfer", files={
                "source": ("notes.txt", good, "text/plain"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("bad extension -> 400", r.status_code == 400, r.text[:200])

            # fake .pptx (random bytes, right name)
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", b"\x00\x01\x02 not a pptx" * 50, "x"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("fake .pptx content -> 400", r.status_code == 400, r.text[:200])

            # empty file
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", b"", "x"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("empty upload -> 400", r.status_code == 400, r.text[:200])

            # oversized: cap enforced on bytes received, not on client-declared size
            main.MAX_UPLOAD_BYTES = 2048
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", good, "x"),   # a real pptx, > 2 KB
                "template": ("t.pptx", tpl, "x"),
            })
            check("oversized upload -> 413 (byte-based cap)", r.status_code == 413, r.text[:200])
            main.MAX_UPLOAD_BYTES = orig_max_bytes

            # excessive slide count
            main.MAX_SLIDES = 3
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", _pptx_bytes(6), "x"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("too many slides -> 413", r.status_code == 413, r.text[:200])
            check("slide-count message names the limit", "maximum supported is 3" in r.text, r.text[:200])
            main.MAX_SLIDES = orig_max_slides

            # happy path: SSE, diagnostics, repeated download
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", _pptx_bytes(4, dense=True), "x"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("valid transfer -> 200 stream", r.status_code == 200, r.text[:200])
            events = _sse_events(r.text)
            done = next((e for e in events if e.get("type") == "done"), None)
            check("SSE produced a done event", done is not None, str(events[-2:]))
            check("SSE progress stages present",
                  {"parsed"} <= {e.get("step") for e in events if e.get("type") == "progress"},
                  str([e.get("step") for e in events]))

            if done:
                check("existing field slides_processed still present", done.get("slides_processed") == 4)
                check("existing field slides_output still present", done.get("slides_output") == 4)
                check("existing field warnings still present", isinstance(done.get("warnings"), list))
                diag = done.get("diagnostics")
                check("diagnostics block present", isinstance(diag, dict), str(done.keys()))
                if isinstance(diag, dict):
                    for k in ("slides_processed", "slides_output", "slides_requiring_review",
                              "warning_count", "error_count", "unsupported_shape_count",
                              "partially_supported_shape_count", "grouped_shape_salvage_count",
                              "overflow_count", "font_reduction_count", "spacing_reduction_count",
                              "transfer_seconds", "validation_seconds", "total_seconds"):
                        check(f"diagnostics has {k}", k in diag, str(diag))
                    check("diagnostics slide counts agree with top-level",
                          diag["slides_processed"] == 4 and diag["slides_output"] == 4, str(diag))
                    check("diagnostics total_seconds >= components",
                          diag["total_seconds"] >= max(diag["transfer_seconds"], diag["validation_seconds"]),
                          str(diag))
                    check("dense deck registered a spacing or font reduction",
                          diag["spacing_reduction_count"] + diag["font_reduction_count"] > 0, str(diag))
                    check("diagnostics leaks no path-like strings",
                          not re.search(r"/tmp|ctm_jobs|\\\\|[A-Za-z]:\\\\", json.dumps(diag)), str(diag))

                url = done["download_url"]
                codes = [client.get(url).status_code for _ in range(3)]
                check("download #1/#2/#3 all 200", codes == [200, 200, 200], str(codes))

            # malformed-but-zip pptx -> parse fails -> sanitized 422 (no traceback/path)
            zb = io.BytesIO()
            with zipfile.ZipFile(zb, "w") as z:
                z.writestr("ppt/presentation.xml", "<not really xml>")
                z.writestr("[Content_Types].xml", "<x/>")
            r = client.post("/api/transfer", files={
                "source": ("deck.pptx", zb.getvalue(), "x"),
                "template": ("t.pptx", tpl, "x"),
            })
            check("corrupt pptx internals -> 422", r.status_code == 422, r.text[:200])
            check("422 message is generic (no traceback/path)",
                  "Traceback" not in r.text and "/tmp" not in r.text and "Errno" not in r.text,
                  r.text[:300])
    finally:
        shutil_rmtree = __import__("shutil").rmtree
        shutil_rmtree(main.TEMP_DIR, ignore_errors=True)
        main.TEMP_DIR = real_tmp
        main.MAX_UPLOAD_BYTES = orig_max_bytes
        main.MAX_SLIDES = orig_max_slides

    print(f"\n{'='*40}\nv2 phase 4: {_passed} passed, {_failed} failed")
    sys.exit(1 if _failed else 0)


if __name__ == "__main__":
    run()
