"""Self-check: source styling that is only *inherited* survives into a template
whose own styles differ (Ch.70 on CUCOM_Presentation_Template regressions):
bullets from the source master, subscripts, Symbol-font glyphs, regular
weight under a bold template body, and a one-slide template's logo.

Run: python test_style_carry.py
"""

import io
import os
import tempfile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from parser import parse_source
from transfer import transfer, _extract_design_images

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
failures = []


def check(name, ok, detail=""):
    print(f"  [{'OK' if ok else 'FAIL'}] {name}" + ("" if ok else f"  {detail}"))
    if not ok:
        failures.append(name)


_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")

td = tempfile.mkdtemp()

# --- template: ONE slide carrying a corner logo; body style is bold ---------
tpl = Presentation()
tpl.slide_width, tpl.slide_height = Inches(13.333), Inches(7.5)
s = tpl.slides.add_slide(tpl.slide_layouts[1])
s.shapes.add_picture(io.BytesIO(_PNG), Inches(10), Inches(0.4), Inches(2.5), Inches(1.2))
body_ph = [p for p in tpl.slide_layouts[1].placeholders if p.placeholder_format.idx == 1][0]
lvl1 = body_ph._element.txBody.find(f"{A}lstStyle")
lvl1.append(etree.fromstring(f'<a:lvl1pPr xmlns:a="{A[1:-1]}"><a:buNone/><a:defRPr b="1"/></a:lvl1pPr>'))
tpl_path = os.path.join(td, "tpl.pptx")
tpl.save(tpl_path)

check("one-slide template: its logo is harvested as decoration",
      len(_extract_design_images(Presentation(tpl_path))) == 1)

# --- source: bullets inherited from master, subscript, Symbol glyph --------
src = Presentation()
src.slide_width, src.slide_height = Inches(13.333), Inches(7.5)
s = src.slides.add_slide(src.slide_layouts[1])
s.shapes.title.text = "Oxygen"
tf = s.placeholders[1].text_frame
p = tf.paragraphs[0]
for text, attrs in (("PO", {}), ("2", {"baseline": "-25000"}), (" rises ", {}), ("", {"sym": 1})):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(20)
    rPr = r._r.get_or_add_rPr()
    if "baseline" in attrs:
        rPr.set("baseline", attrs["baseline"])
    if "sym" in attrs:
        rPr.append(rPr.makeelement(f"{A}sym", {"typeface": "Symbol"}))
src_path = os.path.join(td, "src.pptx")
src.save(src_path)

out = os.path.join(td, "out.pptx")
slides = parse_source(src_path)
transfer(slides, src_path, tpl_path, out)
o = Presentation(out).slides[0]

check("logo stamped on the output slide",
      sum(1 for sh in o.shapes if sh.shape_type == 13) == 1)
body = [sh for sh in o.placeholders if sh.placeholder_format.idx == 1][0]
xml = etree.tostring(body._element).decode()
check("inherited source bullet carried (buChar, not template buNone)", "buChar" in xml, xml[:300])
check("subscript baseline carried", 'baseline="-25000"' in xml)
check("Symbol font carried for U+F0B4", 'typeface="Symbol"' in xml)
runs = [r for para in body.text_frame.paragraphs for r in para.runs]
check("regular-weight runs stay regular under a bold template body",
      all(r.font.bold is False for r in runs), [r.font.bold for r in runs])
check("no empty placeholders left behind",
      all(sh.text_frame.text.strip() for sh in o.placeholders if sh.has_text_frame))

print("\nstyle carry:", "all passed" if not failures else f"{len(failures)} FAILED")
if __name__ == "__main__":
    raise SystemExit(1 if failures else 0)
