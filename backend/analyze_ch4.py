from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

src = Presentation(r'C:\Users\saiva\OneDrive\Desktop\Physiology\Chapter 04 Cellular Organization Cell Functions and Intercellular Connections.pptx')
tpl = Presentation(r'C:\Users\saiva\OneDrive\Desktop\CUCOM_Sports_Yoga_Blank_Template.pptx')

print('=== Template layouts ===')
for i, layout in enumerate(tpl.slide_layouts):
    phs = [(ph.placeholder_format.idx, str(ph.placeholder_format.type)) for ph in layout.placeholders]
    print(f'{i:2d}: {layout.name:35s} phs={phs}')

print()
for idx in [0, 20, 68]:
    slide = src.slides[idx]
    print(f'Source slide {idx+1}: layout={slide.slide_layout.name}')
    for s in slide.shapes:
        if s.has_text_frame:
            ph = s.placeholder_format if s.is_placeholder else None
            phinfo = f'ph_idx={ph.idx} ph_type={ph.type}' if ph else 'no_ph'
            print(f'  text [{phinfo}]: {s.text_frame.text[:80].replace(chr(10), " ")}')
        elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f'  image (PICTURE type)')
        elif s.is_placeholder and s.element.tag.endswith('}pic'):
            print(f'  image (pic placeholder)')
    print()
