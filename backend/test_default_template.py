"""Bundled default template fallback — transfer must work with no template
uploaded and no user-saved one (the state a fresh Render deploy is always in).

Self-contained — no pytest.  Run:  python test_default_template.py
"""
import io
import json
import sys
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
import main

_passed = _failed = 0


def check(name, cond, detail=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK]   {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name}  {detail}")


def _pptx_bytes(n_slides=1):
    prs = Presentation()
    for _ in range(n_slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Title"
        s.placeholders[1].text_frame.text = "point one"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _sse_events(text):
    return [json.loads(l[6:]) for l in text.splitlines() if l.startswith("data: ")]


def run():
    tmp = Path(tempfile.mkdtemp(prefix="deftpl_"))
    saved_orig, meta_orig, default_orig = (
        main.SAVED_TPL_PATH, main.SAVED_TPL_META, main.DEFAULT_TPL_PATH)
    tmp_jobs_orig = main.TEMP_DIR

    # No user-saved template; a bundled default that exists.
    main.SAVED_TPL_PATH = tmp / "nonexistent_saved.pptx"
    main.SAVED_TPL_META = tmp / "nonexistent_saved.json"
    main.DEFAULT_TPL_PATH = tmp / "default_template.pptx"
    main.DEFAULT_TPL_PATH.write_bytes(_pptx_bytes(2))
    main.TEMP_DIR = Path(tempfile.mkdtemp(prefix="deftpl_jobs_"))

    try:
        with TestClient(main.app) as client:
            info = client.get("/api/template/info").json()
            check("info reports the bundled default as available",
                  info.get("saved") is True and info.get("default") is True, str(info))

            # The failing case from the Sep 10 sweep: source only, no template.
            r = client.post("/api/transfer",
                            files={"source": ("deck.pptx", _pptx_bytes(3), "x")})
            check("transfer with no template -> 200 stream", r.status_code == 200, r.text[:200])
            events = _sse_events(r.text)
            done = next((e for e in events if e.get("type") == "done"), None)
            check("SSE produced a done event (default template used)",
                  done is not None, str(events[-2:]))
            check("no error event", not any(e.get("type") == "error" for e in events), r.text[:300])

            # With the default gone, the old 400 still applies.
            main.DEFAULT_TPL_PATH.unlink()
            check("info without any template -> saved false",
                  client.get("/api/template/info").json() == {"saved": False})
            r = client.post("/api/transfer",
                            files={"source": ("deck.pptx", _pptx_bytes(1), "x")})
            check("no template anywhere -> 400", r.status_code == 400, r.text[:200])
    finally:
        main.SAVED_TPL_PATH, main.SAVED_TPL_META, main.DEFAULT_TPL_PATH = (
            saved_orig, meta_orig, default_orig)
        main.TEMP_DIR = tmp_jobs_orig

    print(f"\ndefault template fallback: {_passed} passed, {_failed} failed")
    return _failed == 0


if __name__ == "__main__":
    sys.exit(0 if run() else 1)
