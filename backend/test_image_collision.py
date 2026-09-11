"""Image / body-text collision avoidance (V2).

Root cause: a source picture with no destination picture placeholder lands at
its own (canvas-clamped) source coordinates, with no check against where the
body placeholder actually is. An opaque image can end up on top of body text,
hiding it — a real regression found by rendering Chapter 09 "The Blood"
slide 5 through the CUCOM template.

Self-contained — no pytest.  Run:  python test_image_collision.py
"""
import io
import os
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Emu

from transfer import (
    _rect_intersection_area, _overlap_fraction, _overlaps_meaningfully,
    _adjust_image_for_content_collision, _has_fullbleed_bg, transfer,
)
from parser import parse_source

fails = 0
def check(name, cond, detail=""):
    global fails
    print(f"  [{'OK' if cond else 'FAIL'}] {name}{'' if cond else '  ' + str(detail)}")
    if not cond:
        fails += 1

_PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01'
        b'\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc```\x00\x00'
        b'\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82')

SLIDE_W, SLIDE_H = Inches(13.33), Inches(7.5)


def overlaps(rect, protected):
    return any(_overlaps_meaningfully(rect, p) for p in protected)


# --------------------------------------------------------------------------- #
# Geometry helpers, in isolation
# --------------------------------------------------------------------------- #

check("disjoint rects -> 0 intersection",
      _rect_intersection_area((0, 0, Inches(1), Inches(1)),
                               (Inches(5), Inches(5), Inches(1), Inches(1))) == 0)

check("identical rects -> full overlap fraction 1.0",
      _overlap_fraction((0, 0, Inches(2), Inches(2)), (0, 0, Inches(2), Inches(2))) == 1.0)

# --------------------------------------------------------------------------- #
# Test 1 — image substantially overlaps body -> gets adjusted, clear of it
# --------------------------------------------------------------------------- #
body = (Inches(5), Inches(1), Inches(4), Inches(5))
image = (Inches(4), Inches(3), Inches(6), Inches(1.5))   # deliberately over the body
check("setup: image does start out overlapping body", overlaps(image, [body]))

adjusted, moved, still_colliding = _adjust_image_for_content_collision(
    image, [body], SLIDE_W, SLIDE_H)
check("Test1: collision resolved", not overlaps(adjusted, [body]), adjusted)
check("Test1: image was moved", moved)
check("Test1: not marked still-colliding", not still_colliding)
check("Test1: image stays on-slide",
      adjusted[0] >= 0 and adjusted[1] >= 0
      and adjusted[0] + adjusted[2] <= SLIDE_W and adjusted[1] + adjusted[3] <= SLIDE_H,
      adjusted)
orig_ratio = image[2] / image[3]
new_ratio = adjusted[2] / adjusted[3]
check("Test1: aspect ratio preserved", abs(orig_ratio - new_ratio) < 1e-6,
      (orig_ratio, new_ratio))

# --------------------------------------------------------------------------- #
# Test 2 — image does not overlap body -> unchanged (no over-aggressive moves)
# --------------------------------------------------------------------------- #
far_image = (Inches(0.2), Inches(0.2), Inches(2), Inches(1))
adjusted2, moved2, _ = _adjust_image_for_content_collision(far_image, [body], SLIDE_W, SLIDE_H)
check("Test2: non-overlapping image left untouched", adjusted2 == far_image and not moved2, adjusted2)

# --------------------------------------------------------------------------- #
# Test 3 — image barely touches the body's edge -> below threshold, unchanged
# --------------------------------------------------------------------------- #
# A wide banner whose bottom edge only grazes the top of the body box: most of
# the image and most of the body are untouched by each other.
edge_image = (Inches(4.5), Inches(0.5), Inches(6), Inches(0.6))
frac = _overlap_fraction(edge_image, body)
check("Test3: edge contact is below the collision threshold", frac < 0.15, frac)
adjusted3, moved3, _ = _adjust_image_for_content_collision(edge_image, [body], SLIDE_W, SLIDE_H)
check("Test3: edge-touch image left untouched", adjusted3 == edge_image and not moved3, adjusted3)

# --------------------------------------------------------------------------- #
# Test 4 — image can't move clear without leaving the slide -> scaled instead
# --------------------------------------------------------------------------- #
tiny_slide_w, tiny_slide_h = Inches(6), Inches(4)
full_body = (Inches(0), Inches(0), Inches(6), Inches(4))          # covers the whole slide
big_image = (Inches(1), Inches(1), Inches(4), Inches(2))
adjusted4, moved4, still4 = _adjust_image_for_content_collision(
    big_image, [full_body], tiny_slide_w, tiny_slide_h)
check("Test4: stays fully on-slide even when it can't fully escape",
      adjusted4[0] >= 0 and adjusted4[1] >= 0
      and adjusted4[0] + adjusted4[2] <= tiny_slide_w and adjusted4[1] + adjusted4[3] <= tiny_slide_h,
      adjusted4)
check("Test4: image kept (not dropped)", adjusted4[2] > 0 and adjusted4[3] > 0)
r4 = adjusted4[2] / adjusted4[3]
r4_orig = big_image[2] / big_image[3]
check("Test4: aspect ratio preserved under best-effort scaling",
      abs(r4 - r4_orig) < 1e-6, (r4, r4_orig))

# --------------------------------------------------------------------------- #
# Test 5 — multiple images: later ones avoid both the body and earlier images
# --------------------------------------------------------------------------- #
protected = [body]
placed = []
imgs = [(Inches(4), Inches(3), Inches(4), Inches(1)) for _ in range(3)]  # all start on the body
for im in imgs:
    adj, _, _ = _adjust_image_for_content_collision(im, protected + placed, SLIDE_W, SLIDE_H)
    placed.append(adj)
check("Test5: no image left overlapping the body",
      all(not _overlaps_meaningfully(p, body) for p in placed), placed)
check("Test5: no image left overlapping another placed image",
      all(not _overlaps_meaningfully(placed[i], placed[j])
          for i in range(len(placed)) for j in range(len(placed)) if i != j),
      placed)
check("Test5: all three images kept", len(placed) == 3)


# --------------------------------------------------------------------------- #
# End-to-end helpers
# --------------------------------------------------------------------------- #

def _mk_source(image_rect, slide_w=SLIDE_W, slide_h=SLIDE_H, body_lines=3):
    """Title + body + one free picture at image_rect (left, top, w, h), source coords."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = slide_w, slide_h
    s = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    s.shapes.title.text = "THE BLOOD"
    tf = s.placeholders[1].text_frame
    tf.text = "Blood is defined as liquid connective tissue."
    for _ in range(body_lines - 1):
        tf.add_paragraph().text = "Additional bullet line of body content."
    left, top, w, h = image_rect
    s.shapes.add_picture(io.BytesIO(_PNG), left, top, w, h)
    return prs


def _body_and_image_rects(output_path):
    """Body placeholder rect + transferred *content* image rects — excludes
    the template's own full-slide background/decoration picture, which is
    expected to sit behind everything and isn't a collision candidate."""
    op = Presentation(output_path)
    slide_w, slide_h = op.slide_width, op.slide_height
    slide = op.slides[0]
    body_rect = None
    image_rects = []
    for sh in slide.shapes:
        if sh.shape_type == 13:  # PICTURE
            is_fullbleed = sh.width >= slide_w * 0.98 and sh.height >= slide_h * 0.98
            if not is_fullbleed:
                image_rects.append((sh.left, sh.top, sh.width, sh.height))
        elif sh.is_placeholder and sh.placeholder_format.idx in (1, 2) \
                and sh.has_text_frame and sh.text_frame.text.strip():
            body_rect = (sh.left, sh.top, sh.width, sh.height)
    return body_rect, image_rects


# --------------------------------------------------------------------------- #
# Test 6 — CUCOM: safe-band behavior unchanged, image/body collision resolved,
# no image left inside the protected wedge (the wedge-clamped body rect).
# --------------------------------------------------------------------------- #
default_tpl = Path(__file__).parent / "default_template.pptx"
if default_tpl.exists():
    td = tempfile.mkdtemp()
    # Reproduce the real Ch09-slide-5 geometry: a wide banner image sitting
    # where a right-hand body column would go on a 13.33x7.5in deck.
    src_path = os.path.join(td, "src_cucom.pptx")
    _mk_source((Inches(1.98), Inches(3.22), Inches(9.37), Inches(1.49)),
               body_lines=4).save(src_path)
    out_path = os.path.join(td, "out_cucom.pptx")
    results = transfer(parse_source(src_path), src_path, str(default_tpl), out_path)

    tpl_prs = Presentation(str(default_tpl))
    design_images = __import__("transfer")._extract_design_images(tpl_prs)
    check("Test6: default_template.pptx is still detected as a wedge template",
          _has_fullbleed_bg(design_images, tpl_prs.slide_width, tpl_prs.slide_height))

    body_rect, image_rects = _body_and_image_rects(out_path)
    check("Test6: body placeholder found in output", body_rect is not None)
    if body_rect and image_rects:
        check("Test6: no transferred image overlaps the (wedge-clamped) body",
              all(not _overlaps_meaningfully(r, body_rect) for r in image_rects),
              (body_rect, image_rects))
    check("Test6: no transfer errors", not any(r.errors for r in results),
          [r.errors for r in results])
else:
    print("  [SKIP] Test6/8 — backend/default_template.pptx not present")

# --------------------------------------------------------------------------- #
# Test 7 — non-CUCOM (plain python-pptx template, no full-bleed background):
# no wedge restriction leaks in, but collision avoidance still applies.
# --------------------------------------------------------------------------- #
td2 = tempfile.mkdtemp()
plain_tpl_path = os.path.join(td2, "plain.pptx")
Presentation().save(plain_tpl_path)  # python-pptx's own default template

src2_path = os.path.join(td2, "src_plain.pptx")
_mk_source((Inches(2), Inches(2), Inches(7), Inches(2)), body_lines=4).save(src2_path)
out2_path = os.path.join(td2, "out_plain.pptx")
transfer(parse_source(src2_path), src2_path, plain_tpl_path, out2_path)

plain_prs = Presentation(plain_tpl_path)
plain_design_images = __import__("transfer")._extract_design_images(plain_prs)
check("Test7: plain template is NOT a wedge template",
      not _has_fullbleed_bg(plain_design_images, plain_prs.slide_width, plain_prs.slide_height))

body_rect2, image_rects2 = _body_and_image_rects(out2_path)
if body_rect2 and image_rects2:
    check("Test7: image/body collision still prevented on a plain template",
          all(not _overlaps_meaningfully(r, body_rect2) for r in image_rects2),
          (body_rect2, image_rects2))

# --------------------------------------------------------------------------- #
# Test 8 — the real regression: Ch09-slide-5-shaped input through the actual
# bundled CUCOM template must no longer leave body text hidden behind the
# image (checked structurally here; PNG render is the separate visual check).
# --------------------------------------------------------------------------- #
if default_tpl.exists():
    # Re-use Test 6's output — same exact reproduction of the reported defect.
    body_rect, image_rects = _body_and_image_rects(out_path)
    check("Test8: regression — image no longer overlaps body on the CUCOM template",
          body_rect is not None and image_rects
          and all(not _overlaps_meaningfully(r, body_rect) for r in image_rects),
          (body_rect, image_rects))


print(f"\nimage/body collision: {'all passed' if not fails else f'{fails} FAILED'}")
raise SystemExit(1 if fails else 0)
