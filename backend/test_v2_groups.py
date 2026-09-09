"""Phase 1 — grouped-shape text salvage through parse -> transfer.

Self-contained — no pytest.  Run:  python test_v2_groups.py

Real defect this guards against: a grouped shape containing lecture text was
classified UNSUPPORTED and the whole group (text included) was dropped.
"""
import os
import sys
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))

from parser import parse_source
from transfer import transfer
from validator import validate

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

_passed = _failed = 0


def check(name, cond, detail=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK]   {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name}  {detail}")


def _grp(children_xml: str) -> str:
    return (
        f'<p:grpSp xmlns:p="{_P}" xmlns:a="{_A}">'
        '<p:nvGrpSpPr><p:cNvPr id="30" name="Group 1"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="900000" y="900000"/><a:ext cx="5000000" cy="3000000"/>'
        '<a:chOff x="0" y="0"/><a:chExt cx="5000000" cy="3000000"/></a:xfrm></p:grpSpPr>'
        f'{children_xml}</p:grpSp>'
    )


def _txt_sp(i: int, y: int, text: str) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="TB{i}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="0" y="{y}"/><a:ext cx="5000000" cy="1200000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'
    )


def _rect_sp(i: int) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="R{i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="800000" cy="400000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:sp>'
    )


def _make_deck(path, children_xml, add_title=True):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5 if not add_title else 1])
    if add_title:
        slide.shapes.title.text = "Cell Junctions"
    slide.shapes._spTree.append(etree.fromstring(_grp(children_xml)))
    prs.save(str(path))


def _template(path):
    Presentation().save(str(path))


def _all_text(pptx_path):
    prs = Presentation(pptx_path)
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def run():
    d = Path(tempfile.mkdtemp(prefix="v2_groups_"))

    # --- Group WITH text: salvage into body, PARTIALLY_SUPPORTED, no data loss ---
    src = d / "grp_text.pptx"
    _make_deck(src, _txt_sp(31, 0, "Tight junctions seal adjacent cells")
                    + _txt_sp(32, 1300000, "Gap junctions allow ion flow")
                    + _rect_sp(33))
    slides = parse_source(str(src))
    check("one slide parsed", len(slides) == 1)
    s0 = slides[0]

    body_text = "\n".join(b.full_text for b in s0.body_boxes)
    check("grouped text line 1 salvaged into body", "Tight junctions seal adjacent cells" in body_text, body_text)
    check("grouped text line 2 salvaged into body", "Gap junctions allow ion flow" in body_text, body_text)

    grp_us = [u for u in s0.unsupported_shapes if u.type == "group"]
    check("group recorded as unsupported shape", len(grp_us) == 1, str(s0.unsupported_shapes))
    check("group support_level == PARTIALLY_SUPPORTED",
          grp_us and grp_us[0].support_level == "PARTIALLY_SUPPORTED",
          grp_us[0].support_level if grp_us else "none")

    tpl = d / "tpl.pptx"
    _template(tpl)
    out = d / "grp_text_out.pptx"
    results = transfer(source_slides=slides, source_path=str(src),
                       template_path=str(tpl), output_path=str(out))
    rep = validate(str(out), len(slides), results)
    check("output is valid pptx", rep.ok, str(rep.messages))

    out_text = _all_text(str(out))
    check("salvaged text present in OUTPUT deck", "Gap junctions allow ion flow" in out_text, out_text)

    cw = [c for c in results[0].content_warnings if c["type"] == "group"]
    check("content_warning emitted for group", len(cw) == 1, str(results[0].content_warnings))
    check("content_warning carries support_level",
          cw and cw[0].get("support_level") == "PARTIALLY_SUPPORTED", str(cw))
    check("content_warning message says salvaged, not 'could not be transferred automatically'",
          cw and "salvaged" in cw[0]["message"].lower()
          and "could not be transferred automatically" not in cw[0]["message"].lower(),
          str(cw))
    joined = " ".join(results[0].warnings).lower()
    check("slide warning mentions salvaged grouped text", "salvaged" in joined, joined)

    # --- Group WITHOUT text: SKIPPED_WITH_WARNING, honest skip warning ---
    src2 = d / "grp_empty.pptx"
    _make_deck(src2, _rect_sp(51) + _rect_sp(52), add_title=True)
    slides2 = parse_source(str(src2))
    grp_us2 = [u for u in slides2[0].unsupported_shapes if u.type == "group"]
    check("empty group recorded", len(grp_us2) == 1, str(slides2[0].unsupported_shapes))
    check("empty group support_level == SKIPPED_WITH_WARNING",
          grp_us2 and grp_us2[0].support_level == "SKIPPED_WITH_WARNING",
          grp_us2[0].support_level if grp_us2 else "none")

    out2 = d / "grp_empty_out.pptx"
    results2 = transfer(source_slides=slides2, source_path=str(src2),
                        template_path=str(tpl), output_path=str(out2))
    cw2 = [c for c in results2[0].content_warnings if c["type"] == "group"]
    check("empty-group content_warning uses default 'could not be transferred' phrasing",
          cw2 and "could not be transferred automatically" in cw2[0]["message"], str(cw2))
    check("empty-group slide warning is a skip warning",
          any("skipped" in w.lower() for w in results2[0].warnings), str(results2[0].warnings))

    # --- Nested group: recursion reaches inner text ---
    src3 = d / "grp_nested.pptx"
    inner = _grp(_txt_sp(61, 0, "Desmosomes anchor intermediate filaments"))
    _make_deck(src3, inner + _txt_sp(62, 1500000, "Outer group text too"), add_title=True)
    slides3 = parse_source(str(src3))
    body3 = "\n".join(b.full_text for b in slides3[0].body_boxes)
    check("nested-group inner text salvaged", "Desmosomes anchor intermediate filaments" in body3, body3)
    check("nested-group outer text salvaged", "Outer group text too" in body3, body3)

    print(f"\n{'='*40}\nv2 groups: {_passed} passed, {_failed} failed")
    sys.exit(1 if _failed else 0)


if __name__ == "__main__":
    run()
