"""Validate already-generated Physiology Converted PPTX files (no re-conversion).

Pairs each output with its source by chapter number, then QC-checks:
missing slides, missing text, broken images, blank slides, wrong dimensions,
corrupted ZIP structure, unsupported shapes carried from the source.

Usage (arg order matches batch_convert.py):
    python batch_validate.py [SOURCE_DIR] [TEMPLATE_PPTX] [OUTPUT_DIR]
"""

import re
import sys
import json
import zipfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from pptx import Presentation
from pptx.util import Emu
from parser import parse_source

DESKTOP = Path.home() / "OneDrive" / "Desktop"
SOURCE_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else DESKTOP / "Physiology"
TEMPLATE = Path(sys.argv[2]) if len(sys.argv) > 2 else DESKTOP / "CUCOM Template.pptx"
OUT_DIR = Path(sys.argv[3]) if len(sys.argv) > 3 else DESKTOP / "Physiology Converted"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def chap(name: str):
    m = re.search(r"chapter[\s_]*0*(\d+)", name, re.IGNORECASE)
    return int(m.group(1)) if m else None


def slide_text_len(prs) -> int:
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                n += len(sh.text_frame.text.strip())
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        n += len(c.text_frame.text.strip())
    return n


def broken_images(prs) -> int:
    bad = 0
    for s in prs.slides:
        part = s.part
        for blip in s.element.iter(f"{{{A_NS}}}blip"):
            rid = blip.get(f"{{{R_NS}}}embed")
            if not rid:
                continue
            try:
                blob = part.related_part(rid).blob
                if not blob:
                    bad += 1
            except Exception:
                bad += 1
    return bad


def blank_slides(prs) -> list[int]:
    out = []
    for i, s in enumerate(prs.slides, 1):
        if len(s.shapes) == 0:
            out.append(i)
    return out


def dup_zip_entries(path: Path) -> int:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return len(names) - len(set(names))


def main():
    tpl = Presentation(str(TEMPLATE))
    tpl_dim = (tpl.slide_width, tpl.slide_height)

    src_by_chap = {}
    for p in SOURCE_DIR.glob("*.pptx"):
        c = chap(p.name)
        if c is not None:
            src_by_chap[c] = p

    outputs = sorted(OUT_DIR.glob("*.pptx"), key=lambda p: chap(p.name) or 0)
    report = []

    for out in outputs:
        c = chap(out.name)
        e = {"output": out.name, "source": None, "issues": [], "notes": []}
        try:
            oprs = Presentation(str(out))
        except Exception as ex:
            e["issues"].append(f"Output will not open as PPTX: {ex}")
            report.append(e)
            continue

        o_slides = len(oprs.slides)
        e["slides_output"] = o_slides

        # dimensions
        if (oprs.slide_width, oprs.slide_height) != tpl_dim:
            e["issues"].append(
                f"Slide dimensions {oprs.slide_width}x{oprs.slide_height} "
                f"!= template {tpl_dim[0]}x{tpl_dim[1]}")

        # zip structure
        d = dup_zip_entries(out)
        if d:
            e["issues"].append(f"{d} duplicate ZIP entries (corrupt structure)")

        # blank slides
        bl = blank_slides(oprs)
        if bl:
            e["issues"].append(f"Blank slides (0 shapes): {bl}")

        # broken images
        bi = broken_images(oprs)
        if bi:
            e["issues"].append(f"{bi} unresolved image reference(s)")

        # compare to source
        src = src_by_chap.get(c)
        if src is None:
            e["notes"].append("No matching source chapter found — count check skipped")
        else:
            e["source"] = src.name
            sprs = Presentation(str(src))
            s_slides = len(sprs.slides)
            e["slides_source"] = s_slides
            if o_slides != s_slides:
                e["issues"].append(f"Slide count {o_slides} != source {s_slides}")

            s_text = slide_text_len(sprs)
            o_text = slide_text_len(oprs)
            e["text_source"], e["text_output"] = s_text, o_text
            if s_text > 0 and o_text < 0.85 * s_text:
                e["issues"].append(
                    f"Possible missing text: output {o_text} chars vs source {s_text} "
                    f"({o_text / s_text:.0%})")

            # unsupported shapes carried from source
            try:
                parsed = parse_source(str(src))
                us = {}
                for ps in parsed:
                    for u in ps.unsupported_shapes:
                        us[u.label] = us.get(u.label, 0) + 1
                if us:
                    e["notes"].append(
                        "Source has unsupported shapes (not transferable): "
                        + ", ".join(f"{k}×{v}" for k, v in us.items()))
            except Exception as ex:
                e["notes"].append(f"Source re-parse failed: {ex}")

        report.append(e)

    clean = [e for e in report if not e["issues"]]
    flagged = [e for e in report if e["issues"]]

    print("=" * 70)
    print(f"VALIDATION — {len(report)} generated files")
    print("=" * 70)
    print(f"Clean (no issues)     : {len(clean)}")
    print(f"Flagged with issues   : {len(flagged)}")
    print(f"Template slide size    : {Emu(tpl_dim[0]).inches:.2f} x {Emu(tpl_dim[1]).inches:.2f} in")
    print(f"Output folder         : {OUT_DIR}\n")

    if flagged:
        print("FILES WITH ISSUES:")
        for e in flagged:
            print(f"\n  {e['output']}  (source: {e.get('source', '?')})")
            print(f"    slides: {e.get('slides_output', '?')} output / {e.get('slides_source', '?')} source")
            for iss in e["issues"]:
                print(f"    ISSUE: {iss}")
            for n in e["notes"]:
                print(f"    note:  {n}")

    print("\nNON-BLOCKING NOTES on clean files:")
    any_note = False
    for e in clean:
        if e["notes"]:
            any_note = True
            print(f"  {e['output']}: {'; '.join(e['notes'])}")
    if not any_note:
        print("  (none)")

    (OUT_DIR / "_validation_report.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8")
    print(f"\nFull report: {OUT_DIR / '_validation_report.json'}")


if __name__ == "__main__":
    main()
