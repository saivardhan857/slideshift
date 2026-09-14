"""End-of-slide overlap audit (_find_slide_overlaps).

Root cause this guards against: per-content-type placement code (title vs.
body vs. image) only ever checks a NEW element against what came before it --
nothing looked at the finished slide as a whole. A geometry mismatch between
two placeholders used as parallel text columns (see the OBJECT/BODY
placeholder-role fix, test_placeholder_roles.py) produced exactly this: two
placed elements overlapping, invisible to any single per-element check.

Self-contained — no pytest.  Run:  python test_layout_overlap_qa.py
"""
from pptx import Presentation
from pptx.util import Inches

from transfer import _find_slide_overlaps, _add_background_images

fails = 0
def check(name, cond, detail=""):
    global fails
    print(f"  [{'OK' if cond else 'FAIL'}] {name}{'' if cond else '  ' + str(detail)}")
    if not cond:
        fails += 1

SLIDE_W, SLIDE_H = Inches(13.33), Inches(7.5)

# --------------------------------------------------------------------------- #
# Test 1 — two deliberately overlapping textboxes are caught.
# --------------------------------------------------------------------------- #
prs1 = Presentation()
prs1.slide_width, prs1.slide_height = SLIDE_W, SLIDE_H
s1 = prs1.slides.add_slide(prs1.slide_layouts[6])
tb_a = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
tb_a.text_frame.text = "First overlapping box"
tb_b = s1.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(4), Inches(2))
tb_b.text_frame.text = "Second overlapping box"

conflicts1 = _find_slide_overlaps(s1, SLIDE_W, SLIDE_H)
check("Test1: overlapping textboxes flagged", len(conflicts1) >= 1, conflicts1)

# --------------------------------------------------------------------------- #
# Test 2 — a normal, non-overlapping title+body slide is left alone.
# --------------------------------------------------------------------------- #
prs2 = Presentation()
prs2.slide_width, prs2.slide_height = SLIDE_W, SLIDE_H
s2 = prs2.slides.add_slide(prs2.slide_layouts[1])  # Title and Content
s2.shapes.title.text = "A normal title"
s2.placeholders[1].text_frame.text = "Some body bullet text"

conflicts2 = _find_slide_overlaps(s2, SLIDE_W, SLIDE_H)
check("Test2: normal title+body slide has no false positive", conflicts2 == [], conflicts2)

# --------------------------------------------------------------------------- #
# Test 3 — background/decoration images (added the same way transfer()
# itself adds them, via _add_background_images) are excluded outright, even
# though they sit behind and geometrically overlap everything else. A small
# corner logo counts just as much as a full-bleed background — size doesn't
# matter, only whether it came from _add_background_images.
# --------------------------------------------------------------------------- #
prs3 = Presentation()
prs3.slide_width, prs3.slide_height = SLIDE_W, SLIDE_H
s3 = prs3.slides.add_slide(prs3.slide_layouts[1])
s3.shapes.title.text = "Title over a decorative image"
_PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01'
        b'\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc```\x00\x00'
        b'\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82')
# Deliberately sized/positioned to overlap the title placeholder, to prove
# exclusion is about origin (_add_background_images), not size or position.
bg_ids = _add_background_images(s3, [(_PNG, Inches(0), Inches(0), Inches(4), Inches(2))])

conflicts3 = _find_slide_overlaps(s3, SLIDE_W, SLIDE_H, bg_ids)
check("Test3: decorative image excluded from the audit regardless of size",
      conflicts3 == [], conflicts3)

print(f"\nlayout overlap QA: {'all passed' if not fails else f'{fails} FAILED'}")
if __name__ == "__main__":
    raise SystemExit(1 if fails else 0)
