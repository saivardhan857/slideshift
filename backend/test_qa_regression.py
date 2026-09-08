"""Regression tests from the brutal QA pass (2026-09-08).

Self-contained — no pytest.  Run:  python test_qa_regression.py

Each test reproduces a real defect (or a property that must not regress) found
while stress-testing the production pipeline parse_source -> transfer -> validate.
"""
import io
import sys
import hashlib
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

from parser import parse_source
from transfer import transfer, _BODY_SAFE_RIGHT
from validator import validate
from batch_convert import safe_path, out_name

_EPS = Inches(0.05)  # tolerance for rounding in EMU geometry

TEMPLATE = Path(r"C:\Users\saiva\OneDrive\Desktop\CUCOM Template.pptx")
_passed = _failed = 0


def check(name, cond, detail=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK]   {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name}  {detail}")


def _tmp(name):
    d = Path(tempfile.mkdtemp(prefix="qa_reg_"))
    return d / name


def _two_column_source(path):
    prs = Presentation()  # 10x7.5
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(4), Inches(0.5)).text_frame.text = "Causes"
    left = s.shapes.add_textbox(Inches(0.3), Inches(1), Inches(4), Inches(5)).text_frame
    left.text = "Physiological"
    for i in range(1, 9):
        left.add_paragraph().text = f"{i}. physiological cause number {i} with wording"
    right = s.shapes.add_textbox(Inches(5), Inches(1), Inches(4), Inches(5)).text_frame
    right.text = "Pathological"
    for i in range(1, 13):
        right.add_paragraph().text = f"{i}. pathological cause number {i} with wording"
    prs.save(str(path))
    return path


# ---------------------------------------------------------------------------
# BUG-1  Two-column source must not dump both columns into one placeholder
#        (previously: 2nd column stacked below 1st -> overflowed -> clipped
#         off-slide -> educational content invisible).  Ch49 slides 15/16.
# ---------------------------------------------------------------------------
def test_two_column_not_stacked():
    src = _two_column_source(_tmp("twocol.pptx"))
    slides = parse_source(str(src))
    out = _tmp("twocol_out.pptx")
    transfer(source_slides=slides, source_path=str(src),
             template_path=str(TEMPLATE), output_path=str(out), progress_callback=None)
    op = Presentation(str(out))
    s = op.slides[0]
    phs = {p.placeholder_format.idx: p for p in op.slides[0].placeholders}
    check("two-column slide lands on a layout with a 2nd content placeholder",
          2 in phs, f"placeholder idx present: {sorted(phs)}")
    if 2 in phs:
        t1 = phs[1].text_frame.text.strip()
        t2 = phs[2].text_frame.text.strip()
        check("left content placeholder is non-empty", bool(t1))
        check("right content placeholder is non-empty (2nd column not lost)",
              bool(t2), f"idx2 text={t2!r}")
        check("'Pathological' column text is present somewhere in the two placeholders",
              "pathological cause number 12" in (t1 + " " + t2).lower())


# ---------------------------------------------------------------------------
# BUG-1 safety net: body placeholders get shrink-to-fit autofit so text can
#        never be clipped off-slide even on pathologically dense slides.
# ---------------------------------------------------------------------------
def test_body_autofit_shrink():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Dense"
    body = s.placeholders[1].text_frame
    body.text = "line 1"
    for i in range(80):
        body.add_paragraph().text = f"bullet {i} " * 6
    src = _tmp("dense.pptx")
    prs.save(str(src))
    slides = parse_source(str(src))
    out = _tmp("dense_out.pptx")
    transfer(source_slides=slides, source_path=str(src),
             template_path=str(TEMPLATE), output_path=str(out), progress_callback=None)
    op = Presentation(str(out))
    phs = {p.placeholder_format.idx: p for p in op.slides[0].placeholders}
    if 1 in phs:
        check("body placeholder auto_size == TEXT_TO_FIT_SHAPE",
              phs[1].text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
              f"got {phs[1].text_frame.auto_size}")
        check("dense body text is preserved (not truncated)",
              "bullet 79" in phs[1].text_frame.text)


# ---------------------------------------------------------------------------
# Phase 18  Source files must never be modified by the pipeline.
# ---------------------------------------------------------------------------
def test_source_not_modified():
    src = _two_column_source(_tmp("safe_src.pptx"))
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    slides = parse_source(str(src))
    transfer(source_slides=slides, source_path=str(src), template_path=str(TEMPLATE),
             output_path=str(_tmp("safe_out.pptx")), progress_callback=None)
    after = hashlib.sha256(src.read_bytes()).hexdigest()
    check("source sha256 unchanged after parse+transfer", before == after)


# ---------------------------------------------------------------------------
# Phase 12  One corrupt file in a batch must not stop the others.
# ---------------------------------------------------------------------------
def test_batch_isolation():
    good = _two_column_source(_tmp("g.pptx"))
    d = Path(tempfile.mkdtemp(prefix="qa_iso_"))
    (d / "a.pptx").write_bytes(good.read_bytes())
    (d / "b.pptx").write_bytes(good.read_bytes())
    (d / "bad.pptx").write_bytes(good.read_bytes()[:2000])  # truncated
    (d / "c.pptx").write_bytes(good.read_bytes())
    ok = fail = 0
    for f in sorted(d.glob("*.pptx")):
        try:
            sl = parse_source(str(f))
            transfer(source_slides=sl, source_path=str(f), template_path=str(TEMPLATE),
                     output_path=str(d / (f.stem + "_o.pptx")), progress_callback=None)
            ok += 1
        except Exception:
            fail += 1
    check("3 good files converted despite 1 corrupt", ok == 3 and fail == 1,
          f"ok={ok} fail={fail}")


# ---------------------------------------------------------------------------
# Phase 13  Collision-safe output naming never returns an existing path.
# ---------------------------------------------------------------------------
def test_safe_path_no_overwrite():
    d = Path(tempfile.mkdtemp(prefix="qa_coll_"))
    name = "Physiology_Chapter_01_Converted.pptx"
    (d / name).write_text("existing")
    p2 = safe_path(d, name)
    check("safe_path avoids existing file", p2.name != name and not p2.exists())
    p2.write_text("second")
    p3 = safe_path(d, name)
    check("safe_path avoids two existing files", p3.name not in (name, p2.name))


# ---------------------------------------------------------------------------
# Phase 5  Output slide size is inherited from the template.
# Phase 15 Repeated conversions are structurally identical.
# ---------------------------------------------------------------------------
def test_dims_and_determinism():
    src = _two_column_source(_tmp("det.pptx"))
    tpl_prs = Presentation(str(TEMPLATE))
    tdim = (tpl_prs.slide_width, tpl_prs.slide_height)

    def fp():
        o = _tmp("det_o.pptx")
        sl = parse_source(str(src))
        transfer(source_slides=sl, source_path=str(src), template_path=str(TEMPLATE),
                 output_path=str(o), progress_callback=None)
        p = Presentation(str(o))
        h = hashlib.sha256()
        for s in p.slides:
            for sh in s.shapes:
                h.update(repr((str(sh.shape_type), sh.left, sh.top, sh.width, sh.height,
                               sh.text_frame.text if sh.has_text_frame else "")).encode())
        return (p.slide_width, p.slide_height), h.hexdigest()

    d1, f1 = fp()
    d2, f2 = fp()
    check("output dims == template dims", d1 == tdim, f"{d1} vs {tdim}")
    check("two conversions structurally identical", f1 == f2)


# ---------------------------------------------------------------------------
# BUG-2  Body text must not run across the CUCOM red wedge.  The wedge's left
#        edge inside the body span is ~9.30in; transfer() clamps body content
#        to _BODY_SAFE_RIGHT (9.1in).  Must hold for single- and two-column
#        layouts, must NOT break the BUG-1 two-column fix, must NOT touch
#        title-only slides.
# ---------------------------------------------------------------------------
def _long_body_source(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Mitochondria"
    body = s.placeholders[1].text_frame
    body.text = ("Outer mitochondrial membrane: this forms a continuous envelope "
                 "of the organelle and consists mostly of phospholipids and cholesterol "
                 "and contains a specific membrane protein that forms porin channels.")
    for i in range(6):
        body.add_paragraph().text = (
            f"Point {i}: a deliberately long lecture sentence that would, at full "
            f"placeholder width, wrap across the saturated red wedge on the right.")
    prs.save(str(path))
    return path


def test_bug2_single_column_clear_of_wedge():
    src = _long_body_source(_tmp("longbody.pptx"))
    sl = parse_source(str(src))
    out = _tmp("longbody_out.pptx")
    transfer(source_slides=sl, source_path=str(src), template_path=str(TEMPLATE),
             output_path=str(out), progress_callback=None)
    op = Presentation(str(out))
    phs = {p.placeholder_format.idx: p for p in op.slides[0].placeholders}
    check("single-column body present", 1 in phs)
    if 1 in phs:
        right_edge = phs[1].left + phs[1].width
        check("body right edge <= safe right (clear of red wedge)",
              right_edge <= _BODY_SAFE_RIGHT + _EPS,
              f"right_edge={right_edge/914400:.2f}in  safe={_BODY_SAFE_RIGHT/914400:.2f}in")
        check("full body text preserved after narrowing",
              "porin channels" in phs[1].text_frame.text and "Point 5" in phs[1].text_frame.text)


def test_bug2_two_columns_clear_of_wedge_and_bug1_intact():
    src = _two_column_source(_tmp("tc_wedge.pptx"))
    sl = parse_source(str(src))
    out = _tmp("tc_wedge_out.pptx")
    transfer(source_slides=sl, source_path=str(src), template_path=str(TEMPLATE),
             output_path=str(out), progress_callback=None)
    op = Presentation(str(out))
    phs = {p.placeholder_format.idx: p for p in op.slides[0].placeholders}
    check("two content placeholders present", 1 in phs and 2 in phs)
    if 1 in phs and 2 in phs:
        r1 = phs[1].left + phs[1].width
        r2 = phs[2].left + phs[2].width
        check("left column right edge <= safe right", r1 <= _BODY_SAFE_RIGHT + _EPS,
              f"{r1/914400:.2f}in")
        check("right column right edge <= safe right (no red-wedge overlap)",
              r2 <= _BODY_SAFE_RIGHT + _EPS, f"{r2/914400:.2f}in")
        check("columns do not overlap horizontally", phs[2].left >= r1 - _EPS,
              f"col2.left={phs[2].left/914400:.2f}  col1.right={r1/914400:.2f}")
        # BUG-1 must still hold: both columns carry content, 2nd column not lost
        t1, t2 = phs[1].text_frame.text.lower(), phs[2].text_frame.text.lower()
        check("BUG-1 intact: left column non-empty", bool(t1.strip()))
        check("BUG-1 intact: right column non-empty", bool(t2.strip()))
        check("BUG-1 intact: 'pathological cause number 12' still present",
              "pathological cause number 12" in (t1 + " " + t2))


def test_bug2_title_only_body_geometry_untouched():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Section: Cardiovascular Physiology"
    src = _tmp("titleonly.pptx")
    prs.save(str(src))
    sl = parse_source(str(src))
    out = _tmp("titleonly_out.pptx")
    res = transfer(source_slides=sl, source_path=str(src), template_path=str(TEMPLATE),
                   output_path=str(out), progress_callback=None)
    op = Presentation(str(out))
    s0 = op.slides[0]
    body_phs = [p for p in s0.placeholders if p.placeholder_format.idx in (1, 2)
                and p.has_text_frame and p.text_frame.text.strip()]
    check("title-only slide has no populated body placeholder", not body_phs)
    check("title text preserved", any("Cardiovascular Physiology" in sh.text_frame.text
          for sh in s0.shapes if sh.has_text_frame))
    check("no transfer errors on title-only slide", not any(r.errors for r in res))


if __name__ == "__main__":
    if not TEMPLATE.exists():
        sys.exit(f"template not found: {TEMPLATE}")
    print("QA REGRESSION SUITE")
    print("=" * 50)
    for fn in [test_two_column_not_stacked, test_body_autofit_shrink,
               test_source_not_modified, test_batch_isolation,
               test_safe_path_no_overwrite, test_dims_and_determinism,
               test_bug2_single_column_clear_of_wedge,
               test_bug2_two_columns_clear_of_wedge_and_bug1_intact,
               test_bug2_title_only_body_geometry_untouched]:
        print(f"\n{fn.__name__}:")
        try:
            fn()
        except Exception as e:
            _failed += 1
            import traceback
            print(f"  [ERROR] {fn.__name__}: {e}")
            traceback.print_exc()
    print("\n" + "=" * 50)
    print(f"QA regression: {_passed} passed, {_failed} failed")
    sys.exit(1 if _failed else 0)
