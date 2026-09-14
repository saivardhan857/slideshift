"""Self-contained — run: python test_non_cucom_template.py

Robustness across arbitrary templates:
  * the CUCOM safe-band clamps fire only for a full-bleed-background template;
  * design-image harvesting keeps only decoration repeated across most slides;
  * images/tables placed at source coords are clamped onto a smaller canvas.
"""
import io, tempfile, os
from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from transfer import _has_fullbleed_bg, _extract_design_images, transfer
from parser import parse_source

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

fails = 0
def check(name, cond):
    global fails
    print(f"  [{'OK' if cond else 'FAIL'}] {name}")
    if not cond:
        fails += 1

_PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01'
        b'\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc```\x00\x00'
        b'\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82')

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

fails = 0
def check(name, cond):
    global fails
    print(f"  [{'OK' if cond else 'FAIL'}] {name}")
    if not cond:
        fails += 1

# full-bleed image at origin -> wedge template
check("full-bleed 0,0 -> True",
      _has_fullbleed_bg([(b"", 0, 0, SLIDE_W, SLIDE_H)], SLIDE_W, SLIDE_H))

# CUCOM's is delivered slightly oversized/offset but still covers the slide
check("near-full-bleed -> True",
      _has_fullbleed_bg([(b"", -5000, -5000, SLIDE_W + 20000, SLIDE_H + 20000)],
                        SLIDE_W, SLIDE_H))

# the "compatible" template's inset decoration (L0.82 W11.84) -> not a wedge
check("inset bg 0.82in / 11.84in wide -> False",
      not _has_fullbleed_bg(
          [(b"", Inches(0.82), Inches(0.66), Inches(11.84), Inches(6.35))],
          SLIDE_W, SLIDE_H))

check("no images -> False", not _has_fullbleed_bg([], SLIDE_W, SLIDE_H))

# small logo in a corner -> not a wedge
check("corner logo -> False",
      not _has_fullbleed_bg([(b"", 0, 0, Inches(1.5), Inches(0.8))], SLIDE_W, SLIDE_H))


# --- design-image harvest: repeated decoration kept, per-slide content dropped ---
def _mk_template(path, n_slides, add_repeated_pic, add_unique_pic):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    for i in range(n_slides):
        s = prs.slides.add_slide(blank)
        if add_repeated_pic:
            s.shapes.add_picture(io.BytesIO(_PNG), Inches(0), Inches(0), SLIDE_W, SLIDE_H)
        if add_unique_pic:
            # different bytes per slide -> content, not decoration
            s.shapes.add_picture(io.BytesIO(_PNG + bytes([i % 251])),
                                 Inches(1), Inches(1), Inches(2), Inches(2))
    prs.save(path)

td = tempfile.mkdtemp()
tpl_deco = os.path.join(td, "deco.pptx")
_mk_template(tpl_deco, 6, add_repeated_pic=True, add_unique_pic=True)
di = _extract_design_images(Presentation(tpl_deco))
check("repeated bg harvested, unique photos ignored", len(di) == 1)
check("harvested bg is full-bleed -> wedge=True",
      _has_fullbleed_bg(di, SLIDE_W, SLIDE_H))

tpl_content = os.path.join(td, "content.pptx")
_mk_template(tpl_content, 6, add_repeated_pic=False, add_unique_pic=True)
check("no repeated image -> nothing harvested",
      _extract_design_images(Presentation(tpl_content)) == [])

# --- image + table from a wide source clamped onto a narrow template ---
src_path = os.path.join(td, "src.pptx")
sp = Presentation()
sp.slide_width, sp.slide_height = SLIDE_W, SLIDE_H   # 13.33in wide source
s = sp.slides.add_slide(sp.slide_layouts[5])         # title only
s.shapes.title.text = "Wide slide"
s.shapes.add_picture(io.BytesIO(_PNG), Inches(9), Inches(2), Inches(4), Inches(3))
tb = s.shapes.add_table(2, 2, Inches(8), Inches(5.5), Inches(5), Inches(1.5)).table
tb.cell(0, 0).text = "a"
sp.save(src_path)

narrow_path = os.path.join(td, "narrow.pptx")
np_ = Presentation()
np_.slide_width, np_.slide_height = Inches(10), Inches(7.5)   # 10in template
np_.slides.add_slide(np_.slide_layouts[5])
np_.save(narrow_path)

out_path = os.path.join(td, "out.pptx")
transfer(parse_source(src_path), src_path, narrow_path, out_path)
op = Presentation(out_path)
W, H = op.slide_width, op.slide_height
slack = Emu(int(0.05 * 914400))
overflow = [(sh.shape_type, round(Emu(sh.left).inches, 2), round(Emu(sh.width).inches, 2))
            for sl in op.slides for sh in sl.shapes
            if sh.left is not None and sh.width is not None
            and (sh.left < -slack or sh.left + sh.width > W + slack
                 or sh.top + sh.height > H + slack)]
check(f"no shape off the narrow canvas (got {overflow})", not overflow)

print(f"\nnon-CUCOM template: {'all passed' if not fails else f'{fails} FAILED'}")
if __name__ == "__main__":
    raise SystemExit(1 if fails else 0)
