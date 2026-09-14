"""Placeholder role selection: text vs. picture slots by TYPE, not by idx.

Root cause this guards against: "Content with Caption"-style layouts put a
picture/OBJECT placeholder at idx 1 and the real BODY placeholder at idx 2 --
the opposite of a "Two Content" layout's two idx 1/2 text columns. Code that
assumes idx 1 is always a text column dumps body text into the picture slot,
which has a different top/height than the real body placeholder, producing
overlapping garbled text (and leaves the actual image to float free instead
of filling its intended placeholder).

Self-contained -- no pytest.  Run:  python test_placeholder_roles.py
"""
import io
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from parser import parse_source
from transfer import transfer, _overlaps_meaningfully

fails = 0
def check(name, cond, detail=""):
    global fails
    print(f"  [{'OK' if cond else 'FAIL'}] {name}{'' if cond else '  ' + str(detail)}")
    if not cond:
        fails += 1

_PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01'
        b'\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc```\x00\x00'
        b'\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82')


def _mk_source(tmp_path):
    """Title + two separate body textboxes + one picture -- the same shape
    as a real image-with-caption slide (e.g. an anatomy deck's labeled
    diagram slide)."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    # Both textboxes sit well below the parser's top-of-slide title band
    # (top 20% of a 7.5in slide == 1.5in) so neither is misread as a title --
    # this test targets the body/picture placeholder mix-up, not that heuristic.
    tb1 = s.shapes.add_textbox(Inches(0.5), Inches(2), Inches(4), Inches(1))
    tb1.text_frame.text = "First body fragment"
    tb2 = s.shapes.add_textbox(Inches(0.5), Inches(4), Inches(4), Inches(1))
    tb2.text_frame.text = "Second body fragment"
    s.shapes.add_picture(io.BytesIO(_PNG), Inches(5), Inches(1), Inches(3), Inches(2))
    path = tmp_path / "src.pptx"
    prs.save(path)
    return str(path)


def _mk_template(tmp_path):
    """python-pptx's own default template already ships a 'Content with
    Caption' layout (idx1=OBJECT/picture, idx2=BODY/caption) -- no custom
    template authoring needed to reproduce the bug shape."""
    path = tmp_path / "tpl.pptx"
    Presentation().save(path)
    return str(path)


import tempfile
td = Path(tempfile.mkdtemp())
src_path = _mk_source(td)
tpl_path = _mk_template(td)
out_path = str(td / "out.pptx")

results = transfer(parse_source(src_path), src_path, tpl_path, out_path)
check("no transfer errors", not any(r.errors for r in results), [r.errors for r in results])

out_prs = Presentation(out_path)
slide = out_prs.slides[0]

body_rect = None
object_rect = None
pic_rect = None
for sh in slide.shapes:
    if sh.is_placeholder:
        t = str(sh.placeholder_format.type)
        if t == "BODY (2)":
            body_rect = (sh.left, sh.top, sh.width, sh.height)
        elif t == "OBJECT (7)":
            object_rect = (sh.left, sh.top, sh.width, sh.height)
    if sh.shape_type == 13:  # PICTURE
        pic_rect = (sh.left, sh.top, sh.width, sh.height)

check("BODY placeholder present and holds text",
      body_rect is not None and slide.placeholders and
      any(str(p.placeholder_format.type) == "BODY (2)" and p.text_frame.text.strip()
          for p in slide.placeholders))
check("both body fragments landed in the BODY placeholder (not split into OBJECT)",
      any(str(p.placeholder_format.type) == "BODY (2)"
          and "First body fragment" in p.text_frame.text
          and "Second body fragment" in p.text_frame.text
          for p in slide.placeholders),
      [p.text_frame.text for p in slide.placeholders if p.has_text_frame])
check("OBJECT placeholder was NOT filled with text",
      not any(str(p.placeholder_format.type) == "OBJECT (7)" and p.text_frame.text.strip()
              for p in slide.placeholders if p.has_text_frame))
check("picture landed inside the OBJECT placeholder's frame, not free-floating",
      pic_rect is not None and object_rect is not None
      and abs(pic_rect[0] - object_rect[0]) < Inches(0.05)
      and abs(pic_rect[1] - object_rect[1]) < Inches(0.05),
      (pic_rect, object_rect))
check("body text does not overlap the picture",
      body_rect is not None and pic_rect is not None
      and not _overlaps_meaningfully(body_rect, pic_rect),
      (body_rect, pic_rect))

# --------------------------------------------------------------------------- #
# Sanity: a layout with ONLY an OBJECT placeholder (no sibling BODY) -- the
# common "Title and Content" case -- must still accept body text there.
# --------------------------------------------------------------------------- #
prs2 = Presentation()
s2 = prs2.slides.add_slide(prs2.slide_layouts[1])  # Title and Content
s2.shapes.title.text = "Title"
tpl2_path = str(td / "tpl2.pptx")
prs2.save(tpl2_path)

src2_path = str(td / "src2.pptx")
prs_src2 = Presentation()
ss = prs_src2.slides.add_slide(prs_src2.slide_layouts[6])
tb = ss.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1))
tb.text_frame.text = "Solo body text"
prs_src2.save(src2_path)

results2 = transfer(parse_source(src2_path), src2_path, tpl2_path, str(td / "out2.pptx"))
check("sole-OBJECT layout ('Title and Content') still receives body text",
      not any(r.errors for r in results2))
out2 = Presentation(str(td / "out2.pptx"))
check("text actually landed on the slide",
      any(p.has_text_frame and "Solo body text" in p.text_frame.text
          for p in out2.slides[0].placeholders),
      [p.text_frame.text for p in out2.slides[0].placeholders if p.has_text_frame])

print(f"\nplaceholder roles: {'all passed' if not fails else f'{fails} FAILED'}")
if __name__ == "__main__":
    raise SystemExit(1 if fails else 0)
