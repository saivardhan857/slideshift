"""Parse source PPTX into a structured representation."""

import re
from dataclasses import dataclass, field
from typing import Optional
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
import io

from compat import SupportLevel, classify_shape

_SMARTART_NS = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'
_DRAWING_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


@dataclass
class UnsupportedShape:
    type: str   # "chart" | "smartart" | "group" | "wordart" | "ole" | "media"
    label: str  # human-readable name shown in warnings
    # V2: how much of this shape the transfer engine could carry (compat.SupportLevel value).
    support_level: str = SupportLevel.SKIPPED_WITH_WARNING.value
    detail: str = ""


@dataclass
class TextRun:
    text: str
    bold: Optional[bool] = None
    italic: Optional[bool] = None
    font_size: Optional[float] = None
    font_color: Optional[str] = None  # hex string
    font_name: Optional[str] = None


@dataclass
class Paragraph:
    runs: list[TextRun] = field(default_factory=list)
    level: int = 0  # indent level for bullets

    @property
    def full_text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class TextBox:
    paragraphs: list[Paragraph] = field(default_factory=list)
    placeholder_type: Optional[int] = None  # PP_PLACEHOLDER value
    placeholder_idx: Optional[int] = None
    is_title: bool = False
    is_body: bool = False
    left: int = 0
    top: int = 0
    width: int = 0
    height: int = 0

    @property
    def full_text(self) -> str:
        return "\n".join(p.full_text for p in self.paragraphs)


@dataclass
class ImageData:
    blob: bytes
    content_type: str
    left: int = 0
    top: int = 0
    width: int = 0
    height: int = 0


@dataclass
class TableCell:
    paragraphs: list[Paragraph] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(p.full_text for p in self.paragraphs)


@dataclass
class TableData:
    rows: list[list[TableCell]] = field(default_factory=list)
    left: int = 0
    top: int = 0
    width: int = 0
    height: int = 0

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return len(self.rows[0]) if self.rows else 0


@dataclass
class ParsedSlide:
    index: int  # 0-based
    title: Optional[TextBox] = None
    body_boxes: list[TextBox] = field(default_factory=list)
    images: list[ImageData] = field(default_factory=list)
    tables: list[TableData] = field(default_factory=list)
    layout_name: Optional[str] = None
    unsupported_shapes: list[UnsupportedShape] = field(default_factory=list)
    source_shape_count: int = 0  # total shapes on the source slide

    @property
    def has_title(self) -> bool:
        return self.title is not None and bool(self.title.full_text.strip())

    @property
    def has_body(self) -> bool:
        return any(b.full_text.strip() for b in self.body_boxes)

    @property
    def has_images(self) -> bool:
        return len(self.images) > 0

    @property
    def has_table(self) -> bool:
        return len(self.tables) > 0


_LIST_FRAGMENT_RE = re.compile(r"^\s*(\d+[.)]|[-•*·▪‣●○])\s")


def _looks_like_list_fragment(text: str) -> bool:
    """A bullet/number-prefixed line is body content, not a slide title."""
    return bool(_LIST_FRAGMENT_RE.match(text or ""))


def _parse_color(color) -> Optional[str]:
    try:
        if color and color.type is not None:
            rgb = color.rgb
            return f"{rgb.red:02x}{rgb.green:02x}{rgb.blue:02x}"
    except Exception:
        pass
    return None


def _parse_text_frame(tf) -> list[Paragraph]:
    paragraphs = []
    for para in tf.paragraphs:
        runs = []
        for run in para.runs:
            tr = TextRun(
                text=run.text,
                bold=run.font.bold,
                italic=run.font.italic,
                font_size=run.font.size.pt if run.font.size else None,
                font_color=_parse_color(run.font.color) if run.font.color else None,
                font_name=run.font.name,
            )
            runs.append(tr)
        if not runs and para.text:
            runs.append(TextRun(text=para.text))
        level = para.level or 0
        paragraphs.append(Paragraph(runs=runs, level=level))
    return paragraphs


def _parse_image(shape) -> ImageData:
    img = shape.image
    return ImageData(
        blob=img.blob,
        content_type=img.content_type,
        left=shape.left,
        top=shape.top,
        width=shape.width,
        height=shape.height,
    )


def _parse_pic_placeholder(shape, slide) -> Optional[ImageData]:
    """Extract image from a <p:pic> picture placeholder (shape_type == PLACEHOLDER)."""
    blip = shape.element.find(f'.//{{{_DRAWING_NS}}}blip')
    if blip is None:
        return None
    rId = blip.get(f'{{{_REL_NS}}}embed')
    if not rId:
        return None
    try:
        img_part = slide.part.related_part(rId)
        return ImageData(
            blob=img_part.blob,
            content_type=img_part.content_type,
            left=shape.left,
            top=shape.top,
            width=shape.width,
            height=shape.height,
        )
    except Exception:
        return None


def _parse_table(shape) -> TableData:
    tbl = shape.table
    rows = []
    for row in tbl.rows:
        cells = []
        for cell in row.cells:
            paras = _parse_text_frame(cell.text_frame)
            cells.append(TableCell(paragraphs=paras))
        rows.append(cells)
    return TableData(
        rows=rows,
        left=shape.left,
        top=shape.top,
        width=shape.width,
        height=shape.height,
    )


def _detect_unsupported_shape(shape) -> Optional[UnsupportedShape]:
    """Return an UnsupportedShape (carrying its compat support level) for content
    the transfer engine cannot carry and the user should be told about, else None.

    Fully-supported shapes and decorative shapes with no transferable content
    return None — same warning surface as V1, now routed through compat.py.
    Grouped shapes are handled by _parse_group() in parse_source(), not here.
    """
    c = classify_shape(shape)
    if c.level == SupportLevel.SKIPPED_WITH_WARNING:
        return UnsupportedShape(c.kind, c.label, c.level.value, c.detail)
    return None


def _parse_group(group_shape, parsed) -> bool:
    """Salvage text from a grouped shape into parsed.body_boxes (recursing into
    nested groups). Groups are not transferred structurally; lifting their text
    out is better than dropping the whole group. Records one UnsupportedShape
    for the group whose support level reflects whether text was salvaged.
    Returns True if any text was recovered."""
    salvaged = False
    try:
        children = list(group_shape.shapes)
    except Exception:
        children = []
    for sub in children:
        try:
            st = sub.shape_type
        except Exception:
            st = None
        if st == MSO_SHAPE_TYPE.GROUP:
            if _parse_group(sub, parsed):
                salvaged = True
        else:
            try:
                has_tf = sub.has_text_frame and sub.text_frame.text.strip()
            except Exception:
                has_tf = False
            if has_tf:
                parsed.body_boxes.append(TextBox(
                    paragraphs=_parse_text_frame(sub.text_frame),
                    is_body=True,
                ))
                salvaged = True

    c = classify_shape(group_shape)  # PARTIALLY_SUPPORTED if text present, else SKIPPED_WITH_WARNING
    parsed.unsupported_shapes.append(
        UnsupportedShape(c.kind, c.label, c.level.value, c.detail)
    )
    return salvaged


def parse_source(path: str) -> list[ParsedSlide]:
    prs = Presentation(path)
    slides = []

    for idx, slide in enumerate(prs.slides):
        parsed = ParsedSlide(
            index=idx,
            layout_name=slide.slide_layout.name if slide.slide_layout else None,
            source_shape_count=len(slide.shapes),
        )

        for shape in slide.shapes:
            if shape.has_table:
                parsed.tables.append(_parse_table(shape))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                parsed.images.append(_parse_image(shape))
            elif shape.is_placeholder and shape.element.tag.endswith('}pic'):
                # <p:pic> picture placeholder — shape_type is PLACEHOLDER, not PICTURE
                img = _parse_pic_placeholder(shape, slide)
                if img is not None:
                    parsed.images.append(img)
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                # V2: salvage text from grouped shapes instead of dropping them
                _parse_group(shape, parsed)
            elif shape.has_text_frame:
                ph = shape.placeholder_format if shape.is_placeholder else None
                tb = TextBox(
                    paragraphs=_parse_text_frame(shape.text_frame),
                    placeholder_type=ph.type if ph else None,
                    placeholder_idx=ph.idx if ph else None,
                    left=shape.left,
                    top=shape.top,
                    width=shape.width,
                    height=shape.height,
                )

                # Identify title vs body by placeholder type/idx or position heuristic
                if ph is not None:
                    ph_type_name = ph.type.name if ph.type is not None else ""
                    # Some decks stuff the heading into a mislabelled placeholder
                    # (idx != 0, type OBJECT) that is still *named* "Title N".
                    name_is_title = shape.name.lower().startswith("title")
                    if ph.idx == 0 or ph_type_name in ("TITLE", "CENTER_TITLE"):
                        tb.is_title = True
                    elif name_is_title and not _looks_like_list_fragment(shape.text_frame.text.strip()):
                        tb.is_title = True
                    else:
                        tb.is_body = True
                else:
                    # No placeholder — name beats position/length heuristics.
                    slide_height = prs.slide_height
                    text = shape.text_frame.text.strip()
                    name_is_title = shape.name.lower().startswith('title')
                    if name_is_title and not _looks_like_list_fragment(text):
                        tb.is_title = True
                    elif (len(text) <= 100 and len(shape.text_frame.paragraphs) <= 2
                          and shape.top < slide_height * 0.2
                          and not _looks_like_list_fragment(text)):
                        tb.is_title = True
                    else:
                        tb.is_body = True

                if tb.is_title and parsed.title is None:
                    parsed.title = tb
                else:
                    parsed.body_boxes.append(tb)
            else:
                unhandled = _detect_unsupported_shape(shape)
                if unhandled is not None:
                    parsed.unsupported_shapes.append(unhandled)

        slides.append(parsed)

    return slides
