"""Phase 2 — layout_safety.plan_fit fitting ladder + width sensitivity.

Self-contained — no pytest.  Run:  python test_v2_layout_safety.py
"""
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))

from parser import Paragraph, TextRun, parse_source
from layout_safety import (
    plan_fit, estimate_content_height, demo as ls_demo,
    EMU_PER_IN, MIN_FONT_SCALE, SPACING_REDUCTION_MAX,
)
from transfer import transfer
from validator import validate

_passed = _failed = 0


def check(name, cond, detail=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK]   {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name}  {detail}")


def _para(text, pt=18):
    return Paragraph(runs=[TextRun(text=text, font_size=pt)])


def _paras(n, text, pt=18):
    return [_para(text, pt) for _ in range(n)]


# long line for width-sensitivity (must wrap differently at 3 / 5 / 8 inches)
LINE = ("the quick brown fox jumps over the lazy dog while nine lazy "
        "dogs watch from the shady riverbank nearby today ")  # ~110 chars
# short single-line bullet for the fitting-ladder math (predictable: 1 line each)
SHORT = "Short bullet point here"  # ~23 chars, one line at any sane width


def run():
    ls_demo()  # module self-check

    W = int(6 * EMU_PER_IN)
    H = int(4 * EMU_PER_IN)

    # --- Basic fit ---
    p = plan_fit([_para("Short heading line")], W, H)
    check("short text fits, no adjustment", p.fits and p.action == "none" and p.font_scale == 1.0, p)

    p = plan_fit(_paras(8, SHORT), W, H)
    check("a few normal paragraphs fit untouched", p.fits and p.action == "none", p)

    # --- Width sensitivity: same content, three widths ---
    content = _paras(30, LINE)
    h_wide = estimate_content_height(content, int(8 * EMU_PER_IN))
    h_med = estimate_content_height(content, int(5 * EMU_PER_IN))
    h_narrow = estimate_content_height(content, int(3 * EMU_PER_IN))
    check("narrower box -> taller estimate", h_narrow > h_med > h_wide, (h_narrow, h_med, h_wide))
    check("width actually changes the number materially", h_narrow >= 1.4 * h_wide, (h_narrow, h_wide))

    # --- Slightly overfull -> spacing reduction, font untouched ---
    # 17 single-line bullets in a 4in box: over the 115% trigger, recoverable
    # by a <=20% line-spacing cut alone.
    slight = _paras(17, SHORT)
    p = plan_fit(slight, W, H)
    check("slightly overfull -> reduce_spacing", p.action == "reduce_spacing", p)
    check("slightly overfull keeps intended font size", p.font_scale == 1.0, p)
    check("slightly overfull now fits", p.fits, p)
    check("spacing reduction within safe bound", 0.0 < p.line_spacing_reduction <= SPACING_REDUCTION_MAX, p)

    # --- Denser -> font reduced, but only after spacing is maxed ---
    dense = _paras(20, SHORT)
    p = plan_fit(dense, W, H)
    check("dense -> reduce_font", p.action == "reduce_font", p)
    check("font only reduced after spacing maxed out",
          p.line_spacing_reduction == SPACING_REDUCTION_MAX, p)
    check("dense font scaled below 1.0 but >= floor",
          MIN_FONT_SCALE <= p.font_scale < 1.0, p)
    check("dense now fits", p.fits, p)

    # --- Significantly overfull -> unavoidable overflow, warn, still ship ---
    huge = _paras(80, SHORT)
    p = plan_fit(huge, W, H)
    check("huge -> overflow action", p.action == "overflow", p)
    check("huge does not claim to fit", not p.fits, p)
    check("huge returns a warning", bool(p.warning), p)
    check("huge is floored, never shrinks past MIN_FONT_SCALE", p.font_scale == MIN_FONT_SCALE, p)
    check("huge reports how far it overflows", p.overflow_emu > 0, p)

    # --- Minimum font is a hard floor at every density ---
    for n in (25, 40, 60, 120, 300):
        pp = plan_fit(_paras(n, SHORT), W, H)
        if pp.font_scale < MIN_FONT_SCALE - 1e-9:
            check(f"font floor holds at n={n}", False, pp)
            break
    else:
        check("font never falls below MIN_FONT_SCALE at any density", True)

    # --- Two-column: five length combinations, each column planned on its own box ---
    col_w = int(3.9 * EMU_PER_IN)   # ~equal columns inside the 0.92..9.1in safe band
    col_h = int(5.0 * EMU_PER_IN)
    combos = {
        "short/short": (_paras(3, LINE), _paras(3, LINE)),
        "long/short": (_paras(40, LINE), _paras(3, LINE)),
        "short/long": (_paras(3, LINE), _paras(40, LINE)),
        "long/long": (_paras(40, LINE), _paras(40, LINE)),
        "extreme/extreme": (_paras(120, LINE), _paras(120, LINE)),
    }
    for label, (lft, rgt) in combos.items():
        pl = plan_fit(lft, col_w, col_h)
        pr = plan_fit(rgt, col_w, col_h)
        check(f"two-col [{label}] left column handled (fits or warns)", pl.fits or pl.warning, pl)
        check(f"two-col [{label}] right column handled (fits or warns)", pr.fits or pr.warning, pr)
        check(f"two-col [{label}] left font >= floor", pl.font_scale >= MIN_FONT_SCALE, pl)
        check(f"two-col [{label}] right font >= floor", pr.font_scale >= MIN_FONT_SCALE, pr)

    # --- Integration: dense two-column deck through transfer(), wiring must not crash ---
    d = Path(tempfile.mkdtemp(prefix="v2_ls_"))
    src = d / "dense_twocol.pptx"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(4), Inches(0.5)).text_frame.text = "Causes"
    lt = s.shapes.add_textbox(Inches(0.3), Inches(1), Inches(4), Inches(5)).text_frame
    lt.text = "Left column"
    for i in range(40):
        lt.add_paragraph().text = f"{i}. {LINE}"
    rt = s.shapes.add_textbox(Inches(5), Inches(1), Inches(4), Inches(5)).text_frame
    rt.text = "Right column"
    for i in range(40):
        rt.add_paragraph().text = f"{i}. {LINE}"
    prs.save(str(src))
    tpl = d / "tpl.pptx"
    Presentation().save(str(tpl))

    slides = parse_source(str(src))
    out = d / "out.pptx"
    results = transfer(source_slides=slides, source_path=str(src),
                       template_path=str(tpl), output_path=str(out))
    rep = validate(str(out), len(slides), results)
    check("dense two-col transfer produced a valid deck", rep.ok, str(rep.messages))
    check("dense two-col transfer raised no errors", not any(r.errors for r in results),
          str([r.errors for r in results]))
    op = Presentation(str(out))
    body = " ".join(sh.text_frame.text for sh in op.slides[0].shapes if sh.has_text_frame)
    check("dense two-col content preserved (last line of each column present)",
          "39. " + LINE.strip() in body, body[-200:])

    print(f"\n{'='*40}\nv2 layout_safety: {_passed} passed, {_failed} failed")
    sys.exit(1 if _failed else 0)


if __name__ == "__main__":
    run()
